import os
import sys
import copy
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

# Color palette matching Canva Cyber Theme
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CYAN = RGBColor(0, 210, 255)
COLOR_MUTED = RGBColor(203, 213, 225)
COLOR_RED = RGBColor(255, 75, 75)
COLOR_DARK = RGBColor(11, 25, 44)
COLOR_ACCENT = RGBColor(2, 132, 199)

def update_text(shape, text, font_name="Segoe UI", font_size=None, font_color=COLOR_WHITE, bold=None, align=None):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    orig_align = tf.paragraphs[0].alignment if len(tf.paragraphs) > 0 else None
    
    tf.clear()
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        elif orig_align is not None:
            p.alignment = orig_align
            
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        if font_size is not None:
            run.font.size = font_size
        if font_color is not None:
            run.font.color.rgb = font_color
        if bold is not None:
            run.font.bold = bold

def duplicate_slide(prs, src_index):
    template = prs.slides[src_index]
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)
    for shape in template.shapes:
        new_el = copy.deepcopy(shape.element)
        new_slide.shapes._spTree.append(new_el)
    return new_slide

def move_slide_to_index(prs, from_idx, to_idx):
    slides = prs.slides._sldIdLst
    slide_id = slides[from_idx]
    slides.remove(slide_id)
    slides.insert(to_idx, slide_id)

def main():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    pptx_path = os.path.join(base_dir, 'Blue and White Modern Cyber Security Pitch Deck Presentation.pptx')
    out_copy_path = os.path.join(base_dir, 'Slide_To4_DeTai9_Canva_PitchDeck.pptx')
    
    prs = pptx.Presentation(pptx_path)
    print(f"Loaded {pptx_path} with {len(prs.slides)} slides.")
    
    # Duplicate slide 14 (0-indexed 13) to create Slide 15 for QR Demo
    duplicate_slide(prs, 13)
    print(f"Duplicated slide. Total slides: {len(prs.slides)}") # Now 16 slides
    
    # ----------------------------------------------------
    # SLIDE 1 (Index 0): COVER / BÌA
    # ----------------------------------------------------
    print("Updating Slide 1: Cover...")
    s1 = prs.slides[0]
    update_text(s1.shapes[9], "HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG (PTIT)", font_size=Pt(13), font_color=COLOR_CYAN, bold=True)
    update_text(s1.shapes[10], "BỘ MÔN TƯ DUY HỆ THỐNG • TỔ 4", font_size=Pt(12), font_color=COLOR_MUTED, bold=True)
    update_text(s1.shapes[8], "PHÒNG CHỐNG LỪA ĐẢO", font_size=Pt(36), font_color=COLOR_WHITE, bold=True)
    update_text(s1.shapes[11], "TRỰC TUYẾN", font_size=Pt(48), font_color=COLOR_CYAN, bold=True)
    update_text(s1.shapes[12], "Đề tài 9: Nhận diện hệ thống đối kháng & Tái lập tính nhất thể bằng công nghệ\nNhóm sinh viên thực hiện: Tổ 4 (4 Thành viên)", font_size=Pt(14), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 2 (Index 1): THE PROBLEM / VẤN NẠN HỆ THỐNG
    # ----------------------------------------------------
    print("Updating Slide 2: The Problem...")
    s2 = prs.slides[1]
    update_text(s2.shapes[2], "VẤN NẠN HỆ THỐNG", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s2.shapes[3], "Sinh Viên & Giới Trẻ Trong Tầm Ngắm Tội Phạm Mạng", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s2.shapes[19], "Lừa đảo trực tuyến không còn là hành vi tự phát đơn lẻ mà đã biến chuyển thành các đường dây tội phạm có tổ chức chặt chẽ, trang bị công nghệ cao (Deepfake, AI, botnet), khai thác triệt để khoảng trống nhận thức và áp lực tài chính của sinh viên mới nhập học.", font_size=Pt(12), font_color=COLOR_MUTED)
    update_text(s2.shapes[20], "30.000+ TỶ ĐỒNG THIỆT HẠI\nThiệt hại do lừa đảo trực tuyến tại VN (2024–2026), sinh viên là nhóm dễ tổn thương hàng đầu.", font_size=Pt(11), font_color=COLOR_WHITE)
    update_text(s2.shapes[21], "3 BẪY LỪA ĐẢO PHỔ BIẾN NHẤT\n(1) Việc nhẹ lương cao Shopee/Tiktok • (2) Giả mạo Công an/Thuế • (3) Mạo danh Nhà trường thu học phí.", font_size=Pt(11), font_color=COLOR_WHITE)
    update_text(s2.shapes[22], "ĐẶC ĐIỂM TÂM LÝ DỄ TỔN THƯƠNG\nThiếu kinh nghiệm pháp lý, áp lực tài chính, sợ bị kỷ luật và tâm lý e ngại chia sẻ với gia đình.", font_size=Pt(11), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 3 (Index 2): HỒ SƠ CHỨNG CỨ / 2 CASE STUDY
    # ----------------------------------------------------
    print("Updating Slide 3: Case Evidence...")
    s3 = prs.slides[2]
    update_text(s3.shapes[3], "HỒ SƠ CHỨNG CỨ THỰC TẾ", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s3.shapes[4], "2 Bẫy Lừa Đảo Điển Hình Nhắm Vào Sinh Viên", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s3.shapes[5], "• BẪY 1 - MẠO DANH NHÀ TRƯỜNG THU HỌC PHÍ:\nKẻ gian gửi tin nhắn SMS Brandname mạo danh phòng tài vụ, yêu cầu sinh viên chuyển tiền vào số tài khoản cá nhân mang tên trường để 'bảo lãnh chỉ tiêu', đe dọa xóa tên khỏi danh sách lớp nếu không chuyển tiền trong 24 giờ.\n\n• BẪY 2 - TUYỂN CTV 'VIỆC NHẸ LƯƠNG CAO':\nMời chào làm nhiệm vụ giật đơn sàn Shopee/Tiktok với hoa hồng 10-20%. Hai đơn đầu (50k - 100k) được chuyển trả tiền thật để tạo lòng tin tuyệt đối. Từ đơn thứ 3 bắt đầu đòi nạp tiền triệu rồi chiếm đoạt.", font_size=Pt(11), font_color=COLOR_WHITE)
    
    evidence_img = os.path.join(base_dir, 'assets/images/vn_scam_sms_tuition.jpg')
    if os.path.exists(evidence_img):
        s3.shapes.add_picture(evidence_img, left=10500000, top=1800000, width=6800000, height=7200000)

    # ----------------------------------------------------
    # SLIDE 4 (Index 3): SƠ ĐỒ PHÂN HỆ ĐƯỜNG DÂY LỪA ĐẢO (SƠ ĐỒ 1)
    # ----------------------------------------------------
    print("Updating Slide 4: Scam Subsystem Diagram...")
    s4 = prs.slides[3]
    update_text(s4.shapes[12], "SƠ ĐỒ PHÂN HỆ HỆ THỐNG", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s4.shapes[17], "Đường dây lừa đảo trực tuyến hoạt động có tổ chức chặt chẽ — Không hoạt động tự phát (Phát hiện & nhận diện theo Quyết định 2345/QĐ-NHNN).", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s4.shapes[13], "PHÂN HỆ 1: THU THẬP DỮ LIỆU\nMXH, diễn đàn, dữ liệu rò rỉ ➔ Tiền xử lý, lọc trùng chuẩn hóa tệp sinh viên tiềm năng.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s4.shapes[14], "PHÂN HỆ 2: TẠO LÒNG TIN\nGiả mạo nhà trường, công an; sử dụng công nghệ Deepfake và trả hoa hồng nhỏ mồi nhử ban đầu.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s4.shapes[15], "PHÂN HỆ 3: CHUYỂN TIỀN & RỬA TIỀN\nHướng dẫn chuyển khoản ngân hàng rác, luân chuyển sàn tiền ảo P2P, xóa sạch dấu vết dòng tiền.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s4.shapes[16], "KẾT QUẢ & 4 TÁC ĐỘNG NẶNG NỀ\n(1) Mất tài sản • (2) Ảnh hưởng tâm lý nặng nề • (3) Khó lấy lại tiền • (4) Tăng nguy cơ tái phạm.", font_size=Pt(10), font_color=COLOR_CYAN)
    
    diag1_img = os.path.join(base_dir, 'slides/assets/so_do_phan_he_lua_dao.png')
    if os.path.exists(diag1_img):
        s4.shapes.add_picture(diag1_img, left=6800000, top=4400000, width=4600000, height=3100000)

    # ----------------------------------------------------
    # SLIDE 5 (Index 4 in template): DEMO PHẦN MỀM THỰC TẾ & TÍNH NĂNG
    # ----------------------------------------------------
    print("Updating Slide 5 (Template): Demo & Platform Features...")
    s_demo = prs.slides[4]
    update_text(s_demo.shapes[22], "SẢN PHẨM THỰC TẾ", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_demo.shapes[23], "Cổng Tra Cứu & Trợ Lý AI Lừa Đảo (Tổ 4)", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_demo.shapes[24], "Hệ thống Web App thời gian thực kết hợp Trợ lý AI Client-side giúp sinh viên tự bảo vệ trước mọi tin nhắn đe dọa, đối soát tài khoản học phí chuẩn và gửi báo cáo trực tiếp.", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_demo.shapes[25], "TIỀN KIỂM ĐA NGUỒN (<1S)\nĐối soát SĐT & Whitelist tên miền .edu.vn", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_demo.shapes[26], "TRỢ LÝ AI OFFLINE BÓC TÁCH\nPhân tích tin nhắn đe dọa giục nộp tiền", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_demo.shapes[27], "FIREBASE REALTIME REPORT\nCấp mã Ticket ID tự động theo dõi hồ sơ", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_demo.shapes[28], "TỐI ƯU MOBILE-FIRST\nGiao diện responsive hoàn hảo trên Smartphone", font_size=Pt(10), font_color=COLOR_CYAN)
    
    app_demo_img = os.path.join(base_dir, 'slides/assets/app_portal_preview.png')
    if os.path.exists(app_demo_img):
        s_demo.shapes.add_picture(app_demo_img, left=7200000, top=1200000, width=4200000, height=4500000)

    # ----------------------------------------------------
    # SLIDE 6 (Index 5 in template): SƠ ĐỒ HỆ THỐNG PHÒNG VỆ (SƠ ĐỒ 2)
    # ----------------------------------------------------
    print("Updating Slide 6 (Template): Defense System Diagram & 3 Weaknesses...")
    s_def = prs.slides[5]
    update_text(s_def.shapes[3], "SƠ ĐỒ HỆ THỐNG PHÒNG VỆ", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_def.shapes[4], "6 Chủ Thể Phòng Vệ & 3 Điểm Yếu Liên Kết Trọng Yếu", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_def.shapes[5], "Sức mạnh của hệ thống phòng thủ không phụ thuộc vào số lượng phân hệ, mà quyết định bởi độ bền vững của các mối liên kết giữa 6 chủ thể: Sinh viên, Gia đình, Nhà trường, Nhà mạng, Công an, Ngân hàng.", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_def.shapes[7], "1. Điểm yếu Gia đình ↔ Sinh viên: Khoảng cách địa lý + tâm lý ngại chia sẻ ➔ Khoảng trống thông tin.", font_size=Pt(11), font_color=COLOR_RED)
    update_text(s_def.shapes[8], "2. Điểm yếu Nhà trường ↔ Sinh viên: Kênh 1 chiều, thông tin chính thức đến chậm ➔ Dễ nhầm thông báo thật/giả.", font_size=Pt(11), font_color=COLOR_RED)
    update_text(s_def.shapes[9], "3. Điểm yếu Ngân hàng/Nhà mạng ↔ Công an: Chưa chia sẻ dữ liệu thời gian thực ➔ Chậm phong tỏa dòng tiền.", font_size=Pt(11), font_color=COLOR_RED)
    
    diag2_img = os.path.join(base_dir, 'slides/assets/so_do_phong_ve_diem_yeu.png')
    if os.path.exists(diag2_img):
        # Overlay on right side where Picture 12 was
        s_def.shapes.add_picture(diag2_img, left=10500000, top=1000000, width=5400000, height=8100000)

    # ----------------------------------------------------
    # SLIDE 7 (Index 6 in template): NGUYÊN LÝ TÍNH NHẤT THỂ (WHOLENESS)
    # ----------------------------------------------------
    print("Updating Slide 7 (Template): System Wholeness...")
    s_whole = prs.slides[6]
    update_text(s_whole.shapes[3], "TƯ DUY HỆ THỐNG", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_whole.shapes[4], "Nguyên Lý Tính Nhất Thể: Vì Sao Cảnh Báo Rời Rạc Đều Thất Bại?", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_whole.shapes[5], "Một hệ thống không phải là tập hợp rời rạc của các bộ phận. Khi các mối liên kết bị cắt đứt (Ốc đảo dữ liệu - Data Silos), toàn bộ lá chắn sụp đổ dù từng đơn vị riêng lẻ đều nỗ lực cảnh báo.", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_whole.shapes[12], "1. ỐC ĐẢO DỮ LIỆU (DATA SILOS)\nNgân hàng thấy tiền chuyển, Nhà mạng thấy cuộc gọi, Trường thấy học phí. Không bên nào nắm trọn bức tranh toàn cảnh.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_whole.shapes[13], "2. ĐỘ TRỄ THÔNG TIN TỬ HUYỆT\nKẻ gian tẩu tán tiền trong 3-5 phút, nhưng quy trình công văn hành chính xác minh phong tỏa tài khoản mất 24h đến vài ngày.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_whole.shapes[14], "3. TÁI LẬP TÍNH NHẤT THỂ\nCần mạng lưới liên kết nhất thể đa tác nhân: Sinh viên tra cứu - AI phân tích - Công an xử lý - Gia đình bảo bọc.", font_size=Pt(10), font_color=COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 8 (Index 7 in template): VÒNG LẶP PHẢN HỒI & ĐIỂM ĐÒN BẨY
    # ----------------------------------------------------
    print("Updating Slide 8 (Template): Feedback Loops & Leverage Points...")
    s_fb = prs.slides[7]
    update_text(s_fb.shapes[1], "VÒNG PHẢN HỒI & ĐIỂM ĐÒN BẨY", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_fb.shapes[2], "Causal Loops (R1, R2) & Điểm Can Thiệp Meadows", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_fb.shapes[3], "• VÒNG LẶP KHUẾCH ĐẠI R1 (LÒNG TIN GIẢ MẠO):\nHoa hồng nhỏ ban đầu (50k - 100k) ➔ Tăng niềm tin ➔ Nạp số tiền lớn hơn ➔ Bị chiếm đoạt và đe dọa.\n\n• VÒNG LẶP R2 (MÙA VỤ & HOẢNG LOẠN):\nĐầu kỳ thu học phí ➔ Tâm lý sợ bị xóa môn/hủy học ➔ Vội vàng chuyển khoản vào STK cá nhân giả mạo mà không kịp đối soát.\n\n• ĐIỂM ĐÒN BẨY HỆ THỐNG CỦA DONELLA MEADOWS:\nCan thiệp cấu trúc luồng thông tin (Information Flows): Cung cấp công cụ tiền kiểm tức thời ngay tại thời điểm nhận tin nhắn nghi vấn, chặn đứng dòng tiền trước khi rời khỏi tài khoản sinh viên.", font_size=Pt(11), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 9 (Index 8 in template): CƠ ĐỒ TƯ DUY KIẾN TRÚC HỆ THỐNG (MINDMAP)
    # ----------------------------------------------------
    print("Updating Slide 9 (Template): Comprehensive Mindmap...")
    s_mm = prs.slides[8]
    update_text(s_mm.shapes[2], "CƠ ĐỒ TƯ DUY HỆ THỐNG", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_mm.shapes[3], "Mạng Lưới Phòng Vệ & Kiến Trúc Tác Nhân Liên Kết", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_mm.shapes[4], "Sơ đồ kiến trúc liên kết phân tầng từ Nhận thức -> Tiền kiểm công nghệ -> Báo cáo thời gian thực -> Can thiệp pháp lý.", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_mm.shapes[8], "TẦNG 1: SINH VIÊN\nNâng cao nhận thức, thói quen đối soát thông tin trước khi chuyển tiền.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_mm.shapes[11], "TẦNG 2: NỀN TẢNG SỐ\nCổng tra cứu đa nguồn <1s & Trợ lý AI bóc tách kịch bản thao túng.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_mm.shapes[14], "TẦNG 3: NHÀ TRƯỜNG & GIA ĐÌNH\nKênh xác thực học phí 24/7 & Chỗ dựa tinh thần, tài chính vững chắc.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_mm.shapes[17], "TẦNG 4: CÔNG AN & NGÂN HÀNG\nCơ chế phong tỏa dòng tiền khẩn cấp & Điều tra truy vết đường dây.", font_size=Pt(10), font_color=COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 10 (Index 9 in template): KIẾN TRÚC GIẢI PHÁP 3 LỚP PHÒNG VỆ
    # ----------------------------------------------------
    print("Updating Slide 10 (Template): Solution 3-Layer Architecture...")
    s_arch = prs.slides[9]
    update_text(s_arch.shapes[2], "KIẾN TRÚC GIẢI PHÁP", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_arch.shapes[6], "3 Lớp Phòng Vệ Của Cổng Tra Cứu Tổ 4", font_size=Pt(24), font_color=COLOR_WHITE, bold=True)
    update_text(s_arch.shapes[7], "Lớp 1: Tiền kiểm đa nguồn (SĐT, Email edu, STK theo blacklist Bộ Công An)", font_size=Pt(11), font_color=COLOR_WHITE)
    update_text(s_arch.shapes[3], "Lớp 2: Trợ lý AI Offline bóc tách kịch bản thao túng (Local AI bảo mật)", font_size=Pt(11), font_color=COLOR_WHITE)
    update_text(s_arch.shapes[4], "Lớp 3: Tiếp nhận & Cấp mã phản ánh Realtime (Firebase Cloud)", font_size=Pt(11), font_color=COLOR_WHITE)
    update_text(s_arch.shapes[5], "Cộng đồng: Đồng bộ cảnh báo tức thời, kết nối mạng lưới phòng thủ chung", font_size=Pt(11), font_color=COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 11 (Index 10 in template): YÊU CẦU CHỨC NĂNG & PHI CHỨC NĂNG
    # ----------------------------------------------------
    print("Updating Slide 11 (Template): FR & NFR...")
    s_fr = prs.slides[10]
    update_text(s_fr.shapes[3], "YÊU CẦU CHỨC NĂNG (FUNCTIONAL REQUIREMENTS - FR)", font_size=Pt(18), font_color=COLOR_CYAN, bold=True)
    update_text(s_fr.shapes[4], "• FR-1 (Tra cứu đa nguồn SĐT, Email, STK): Đối chiếu kho blacklist quốc gia; kiểm tra định dạng tên miền chính thống .edu.vn và đối soát tài khoản học phí chuẩn.\n• FR-2 (Phân tích kịch bản bằng AI Offline): Bóc tách các mẫu câu đe dọa, giục chuyển tiền gấp, mạo danh cơ quan chức năng, cảnh báo mức độ rủi ro kèm giải thích chi tiết.\n• FR-3 (Hệ thống tiếp nhận phản ánh): Thu thập dữ liệu báo cáo có phân loại, tự động cấp mã Ticket ID (HS-TDHT-XXXX) để theo dõi và đồng bộ lên Firebase.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_fr.shapes[5], "YÊU CẦU PHI CHỨC NĂNG (NON-FUNCTIONAL REQUIREMENTS - NFR)", font_size=Pt(18), font_color=COLOR_CYAN, bold=True)
    update_text(s_fr.shapes[6], "• NFR-1 (Tốc độ & Hiệu năng): Thời gian phản hồi tra cứu < 1 giây; bộ nhớ tải trang nhẹ, hoạt động mượt mà ngay cả khi kết nối 3G/4G yếu.\n• NFR-2 (Bảo mật & Quyền riêng tư): AI chạy 100% Client-side (Offline), nội dung tin nhắn không bao giờ tải lên máy chủ ngoài máy người dùng.\n• NFR-3 (Khả năng sử dụng & Tương thích): Thiết kế chuẩn Mobile-first, hỗ trợ Dark Mode chống lóa, tương thích 100% với trình duyệt trên điện thoại.", font_size=Pt(10), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 12 (Index 11 in template): MA TRẬN ĐÁNH GIÁ RỦI RO & KIỂM SOÁT
    # ----------------------------------------------------
    print("Updating Slide 12 (Template): Risk Matrix & Mitigations...")
    s_risk = prs.slides[11]
    update_text(s_risk.shapes[3], "RỦI RO VẬN HÀNH & KỸ THUẬT", font_size=Pt(18), font_color=COLOR_RED, bold=True)
    update_text(s_risk.shapes[4], "• RỦI RO 1 - BÁO CÁO ẢO & SPAM DỮ LIỆU (Mức độ: Cao):\nKẻ xấu cố tình gửi số điện thoại giả hoặc spam tin nhắn làm loãng dữ liệu cộng đồng.\n➔ Kiểm soát: Áp dụng cơ chế Rate-limiting (giới hạn tần suất gửi) và thuật toán lọc trùng tự động bằng hàm băm (Hash Fingerprint).\n\n• RỦI RO 2 - LỆCH CHUẨN AI (FALSE POSITIVE / FALSE NEGATIVE) (Mức độ: Trung bình):\nBỏ sót thủ đoạn lừa đảo mới hoặc nhận nhầm tin nhắn nhắc nộp tiền thật của nhà trường.\n➔ Kiểm soát: Kết hợp phân tích luật ngữ nghĩa (Rule-based) với mô hình AI cục bộ; luôn dán nhãn mức độ tin cậy kèm bằng chứng.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_risk.shapes[5], "RỦI RO HẠ TẦNG & PHÁP LÝ", font_size=Pt(18), font_color=COLOR_CYAN, bold=True)
    update_text(s_risk.shapes[6], "• RỦI RO 3 - NGHẼN TẢI CAO ĐIỂM ĐẦU NĂM HỌC (Mức độ: Trung bình):\nLượng sinh viên tra cứu tăng đột biến khi có đợt đóng học phí làm chậm hệ thống.\n➔ Kiểm soát: Thiết kế kiến trúc phi tập trung, bộ đệm dữ liệu cục bộ (IndexedDB/Cache Storage), CDN phân tán.\n\n• RỦI RO 4 - TRÁCH NHIỆM PHÁP LÝ & TÍNH KHÁCH QUAN (Mức độ: Thấp):\nNguy cơ tranh chấp khi gắn nhãn cảnh báo một số điện thoại hoặc tài khoản.\n➔ Kiểm soát: Dán nhãn rõ 'Công cụ tham khảo hỗ trợ sinh viên', trích dẫn nguồn xác thực cơ quan chức năng, có quy trình khiếu nại.", font_size=Pt(10), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 13 (Index 12 in template): ĐỘI NGŨ THỰC HIỆN DỰ ÁN (TỔ 4)
    # ----------------------------------------------------
    print("Updating Slide 13 (Template): Team Organization...")
    s_team = prs.slides[12]
    update_text(s_team.shapes[25], "ĐỘI NGŨ DỰ ÁN (TỔ 4)", font_size=Pt(24), font_color=COLOR_CYAN, bold=True)
    update_text(s_team.shapes[26], "Nhóm sinh viên thực hiện Đề tài 9 - Học phần Tư duy hệ thống (PTIT). 4 thành viên phối hợp chặt chẽ đảm bảo tính học thuật và sản phẩm công nghệ ứng dụng thực tế.\nThành viên 4: Lê Đình Vũ (UI/UX Designer & Presentation Lead - Thiết kế giao diện Web App, đồ họa Canva & tài liệu).", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_team.shapes[19], "Vũ Đức Trường", font_size=Pt(14), font_color=COLOR_CYAN, bold=True)
    update_text(s_team.shapes[20], "Fullstack & AI Engineer\nPhát triển Cổng tra cứu, Trợ lý AI Offline & Firestore Realtime", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_team.shapes[21], "Trịnh Đình Dũng", font_size=Pt(14), font_color=COLOR_CYAN, bold=True)
    update_text(s_team.shapes[22], "System Architect\nMô hình hóa hệ thống, Causal Loops & Phân tích tính nhất thể", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_team.shapes[23], "Nguyễn Duy Hoàng", font_size=Pt(14), font_color=COLOR_CYAN, bold=True)
    update_text(s_team.shapes[24], "Domain Researcher\nNghiên cứu nghiệp vụ, khảo sát thực tế & thu thập dữ liệu", font_size=Pt(10), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # SLIDE 14 (Index 13 in template): QUY TRÌNH ỨNG PHÓ KHỦNG HOẢNG
    # ----------------------------------------------------
    print("Updating Slide 14 (Template): Crisis Response Procedure...")
    s_crisis = prs.slides[13]
    update_text(s_crisis.shapes[10], "BIỆN PHÁP CẦN LÀM NGAY KHI ĐÃ BỊ LỪA", font_size=Pt(22), font_color=COLOR_RED, bold=True)
    update_text(s_crisis.shapes[14], "Khi phát hiện đã lỡ chuyển tiền hoặc cài đặt app lạ, sinh viên cần bình tĩnh thực hiện ngay 4 bước khẩn cấp theo thứ tự ưu tiên:", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_crisis.shapes[13], "BƯỚC 1: CẮT KẾT NỐI & NGẮT APP ĐỘC HẠI\nBật Chế độ máy bay, tắt Wifi/4G. Gỡ ngay file APK/app lạ (eTax, VNeID giả). Tháo SIM nếu cần để ngăn chiếm quyền trợ năng.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_crisis.shapes[11], "BƯỚC 2: PHONG TỎA TÀI KHOẢN NGÂN HÀNG\nGọi ngay Hotline ngân hàng yêu cầu khóa thẻ và tài khoản khẩn cấp, hoặc thao tác khóa tạm thời trên Mobile Banking từ máy khác.", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_crisis.shapes[12], "BƯỚC 3 & 4: BÁO CƠ QUAN CHỨC NĂNG & CHIA SẺ GIA ĐÌNH\nChụp chứng cứ giao dịch, gọi 156 / báo Công an. Chia sẻ ngay với gia đình để được bảo vệ, không tin dịch vụ 'thu hồi tiền' lừa đảo.", font_size=Pt(10), font_color=COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 15 (Index 15 - Duplicated Slide): TRẢI NGHIỆM TRỰC TUYẾN & QUÉT MÃ QR
    # ----------------------------------------------------
    print("Updating Slide 15: Online Experience & QR Code...")
    s_qr = prs.slides[15]
    update_text(s_qr.shapes[10], "TRẢI NGHIỆM CỔNG TRA CỨU & MÃ QR", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_qr.shapes[14], "Cổng tra cứu trực tuyến và Trợ lý AI bóc tách tin nhắn lừa đảo của Tổ 4 đã sẵn sàng phục vụ sinh viên hoàn toàn miễn phí. Quét mã QR bằng điện thoại để trải nghiệm ngay:", font_size=Pt(11), font_color=COLOR_MUTED)
    update_text(s_qr.shapes[11], "🌐 ĐƯỜNG DẪN ỨNG DỤNG:\nhttps://ductruongvu23.github.io/chude9to4/app/index.html", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_qr.shapes[13], "📞 TỔNG ĐÀI QUỐC GIA 156:\nTiếp nhận phản ánh cuộc gọi & tin nhắn rác / lừa đảo", font_size=Pt(10), font_color=COLOR_WHITE)
    update_text(s_qr.shapes[12], "🛡️ KÊNH CHÍNH THỐNG:\nchongluadao.vn • canhsat.gov.vn • ais.gov.vn", font_size=Pt(10), font_color=COLOR_CYAN)
    
    qr_img = os.path.join(base_dir, 'slides/assets/qr_web_app.png')
    if os.path.exists(qr_img):
        s_qr.shapes.add_picture(qr_img, left=10500000, top=2000000, width=5400000, height=5400000)

    # ----------------------------------------------------
    # SLIDE 16 (Index 14 originally - Thank You): TUYÊN NGÔN HỆ THỐNG & CẢM ƠN
    # ----------------------------------------------------
    print("Updating Slide 16: System Manifesto & Thank You...")
    s_ty = prs.slides[14]
    update_text(s_ty.shapes[9], "HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG (PTIT)", font_size=Pt(13), font_color=COLOR_CYAN, bold=True)
    update_text(s_ty.shapes[10], "BỘ MÔN TƯ DUY HỆ THỐNG • TỔ 4", font_size=Pt(12), font_color=COLOR_MUTED, bold=True)
    update_text(s_ty.shapes[8], "TUYÊN NGÔN TƯ DUY HỆ THỐNG", font_size=Pt(22), font_color=COLOR_CYAN, bold=True)
    update_text(s_ty.shapes[11], "CẢM ƠN THẦY CÔ & CÁC BẠN!", font_size=Pt(36), font_color=COLOR_WHITE, bold=True)
    update_text(s_ty.shapes[12], "\"Không có giải pháp công nghệ nào bảo vệ được sinh viên nếu các mối liên kết con người bị cắt đứt. Sức mạnh phòng thủ tối thượng nằm ở sự kết nối nhất thể giữa Sinh viên — Gia đình — Nhà trường — Xã hội.\"\n\n🛡️ 3 KHÔNG: Không vội chuyển tiền • Không click link lạ • Không cài app ngoài Store.\n✅ 3 NÊN: Nên tra cứu Cổng Tổ 4 • Nên chia sẻ với gia đình • Nên báo Công an / 156 ngay.", font_size=Pt(12), font_color=COLOR_WHITE)

    # ----------------------------------------------------
    # REORDER SLIDES TO LOGICAL PRESENTATION ORDER (1 TO 16)
    # ----------------------------------------------------
    print("Reordering slides into optimal presentation narrative flow...")
    # Current slide positions:
    # 0: Cover
    # 1: Problem
    # 2: Case Evidence
    # 3: Scam Subsystem Diagram (Sơ đồ 1)
    # 4: Demo Features
    # 5: Defense System Diagram (Sơ đồ 2)
    # 6: Wholeness
    # 7: Feedback Loops
    # 8: Mindmap
    # 9: 3-Layer Solution
    # 10: FR & NFR
    # 11: Risk Matrix
    # 12: Team
    # 13: Crisis Response
    # 14: Thank You
    # 15: QR & App Demo
    
    # We want:
    # 1. Cover (0)
    # 2. Problem (1)
    # 3. Case Evidence (2)
    # 4. Scam Subsystem Diagram (3)
    # 5. Feedback Loops (7)
    # 6. Defense System Diagram (5)
    # 7. Wholeness (6)
    # 8. Mindmap (8)
    # 9. 3-Layer Solution (9)
    # 10. Demo Features (4)
    # 11. FR & NFR (10)
    # 12. Risk Matrix (11)
    # 13. Crisis Response (13)
    # 14. Team (12)
    # 15. QR & App Demo (15)
    # 16. Thank You (14)
    
    # Let's reorder:
    # Move Feedback Loops (idx 7) to idx 4 (right after Scam Subsystems)
    move_slide_to_index(prs, 7, 4)
    # Current: [0, 1, 2, 3, 7(now 4), 4(now 5), 5(now 6), 6(now 7), 8, 9, 10, 11, 12, 13, 14, 15]
    # Move Demo Features (was 4, now 5) to position after 3-Layer Solution (after 9)
    # Let's find index of Demo features and move it:
    move_slide_to_index(prs, 5, 9)
    
    # Move Crisis Response (now idx 13) before Team (now idx 12)
    move_slide_to_index(prs, 13, 12)
    
    # Move QR Demo (now idx 15) before Thank You (now idx 14)
    # Let's check: QR Demo was 15, Thank you was 14.
    move_slide_to_index(prs, 15, 14)
    
    # Save to both paths
    prs.save(pptx_path)
    prs.save(out_copy_path)
    print(f"\nSUCCESS: Presentation updated and saved to:")
    print(f"  1. {pptx_path} ({os.path.getsize(pptx_path):,} bytes)")
    print(f"  2. {out_copy_path} ({os.path.getsize(out_copy_path):,} bytes)")

if __name__ == '__main__':
    main()
