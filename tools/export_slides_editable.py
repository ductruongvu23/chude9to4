"""Xuất slides/index.html ra PPTX CHỈNH SỬA ĐƯỢC (dùng để nhập vào Canva / PowerPoint).

Khác với tools/export_slides.py (mỗi slide là 1 ảnh chụp → không sửa được), file này dựng:
  • Ảnh nền mỗi slide: nền, khung thẻ, đường viền, emoji (không chứa chữ, không chứa ảnh minh họa)
  • Hộp chữ thật (font Open Sans / Roboto Mono có sẵn trên Canva, đúng cỡ, màu, đậm/nghiêng, căn lề)
  • Ảnh minh họa tách riêng (thay/di chuyển được), giữ cắt khung và bo góc
  • Ghi chú người nói

Cách dùng (từ thư mục gốc dự án):
    python tools/export_slides_editable.py

Kết quả: slides/Slide_To4_DeTai9_ChinhSua.pptx
Yêu cầu: Microsoft Edge, python-pptx, Pillow.
"""

import html
import io
import json
import math
import os
import re
import sys
import tempfile
from urllib.parse import urlparse, unquote

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(__file__))
from export_slides import (APP_URL, ROOT, SLIDE_H, SLIDE_W, SLIDES_DIR,  # noqa: E402
                           find_browser, read_slides, run_browser, start_server)

OUT_PPTX = os.path.join(SLIDES_DIR, 'Slide_To4_DeTai9_ChinhSua.pptx')
WRAPPER_NAME = '__export_editable.html'   # trang tạm, xóa sau khi chạy
BG_SCALE = 2

EMU_PER_PX = 12192000 / SLIDE_W           # 7620 EMU / px (slide 13.333 in = 1600 px)
PT_PER_PX = 960 / SLIDE_W                 # 0.6 pt / px

# Phải trùng với FONT_SANS / FONT_MONO trong tools/slide_extract.js (font có sẵn trên Canva)
FONT_SANS = 'Open Sans'
FONT_MONO = 'Roboto Mono'

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass


def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def wrapper_html(mode, slide):
    return f"""<!doctype html><html><head><meta charset="utf-8">
<style>html,body{{margin:0;padding:0;overflow:hidden;background:#000}}
iframe{{border:0;width:{SLIDE_W}px;height:{SLIDE_H}px;display:block}}</style></head>
<body><iframe id="slideFrame"></iframe><pre id="out"></pre>
<script src="tools/slide_extract.js"></script></body></html>"""


def extract_slide(browser, port, profile_dir, num):
    url = f'http://127.0.0.1:{port}/{WRAPPER_NAME}?mode=extract&slide={num}'
    with tempfile.NamedTemporaryFile(suffix='.html', delete=False) as tmp:
        out_path = tmp.name
    try:
        import subprocess
        cmd = [browser, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--no-first-run',
               f'--user-data-dir={profile_dir}', '--virtual-time-budget=12000',
               f'--window-size={SLIDE_W},{SLIDE_H}', '--dump-dom', url]
        res = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=180)
        m = re.search(r'<pre id="out">(.*?)</pre>', res.stdout, re.S)
        if not m or not m.group(1).strip():
            sys.exit(f'Không đo được bố cục slide {num}.')
        return json.loads(html.unescape(m.group(1)))
    finally:
        os.unlink(out_path)


def screenshot_background(browser, port, profile_dir, work_dir, num):
    png = os.path.join(work_dir, f'bg_{num:02d}.png')
    url = f'http://127.0.0.1:{port}/{WRAPPER_NAME}?mode=bg&slide={num}'
    run_browser(browser, profile_dir, [
        f'--window-size={SLIDE_W},{SLIDE_H}',
        f'--force-device-scale-factor={BG_SCALE}',
        f'--screenshot={png}',
        url,
    ])
    if not os.path.exists(png):
        sys.exit(f'Chụp nền slide {num} thất bại.')
    img = Image.open(png).convert('RGB')
    expected = (SLIDE_W * BG_SCALE, SLIDE_H * BG_SCALE)
    if img.size != expected:
        img = img.crop((0, 0, min(img.width, expected[0]), min(img.height, expected[1]))).resize(expected, Image.LANCZOS)
    jpg = os.path.join(work_dir, f'bg_{num:02d}.jpg')
    img.save(jpg, 'JPEG', quality=90, subsampling=0, optimize=True)
    return jpg


# ------------------------------------------------------------------ ảnh minh họa
def local_path(src):
    path = unquote(urlparse(src).path).lstrip('/')
    return os.path.join(ROOT, *path.split('/'))


def prepare_image(info, work_dir, idx):
    """Cắt ảnh theo object-fit: cover và bo góc giống trên web, trả về file tạm."""
    img = Image.open(local_path(info['src'])).convert('RGBA')
    box_w, box_h = info['w'], info['h']
    if info.get('objectFit') == 'cover':
        target = box_w / box_h
        cur = img.width / img.height
        if cur > target:
            new_w = int(img.height * target)
            left = (img.width - new_w) // 2
            img = img.crop((left, 0, left + new_w, img.height))
        else:
            new_h = int(img.width / target)
            top = (img.height - new_h) // 2
            img = img.crop((0, top, img.width, top + new_h))
    radius = info.get('radius') or 0
    out = os.path.join(work_dir, f'img_{idx}.png')
    if radius > 0:
        scale = img.width / box_w
        mask = Image.new('L', img.size, 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, img.width - 1, img.height - 1),
                                               radius=int(radius * scale), fill=255)
        img.putalpha(mask)
        img.save(out, 'PNG', optimize=True)
    else:
        out = out[:-4] + '.jpg'
        img.convert('RGB').save(out, 'JPEG', quality=90)
    return out


# ------------------------------------------------------------------ hộp chữ
def font_for(family):
    fam = (family or '').lower()
    return FONT_MONO if ('mono' in fam or 'consolas' in fam or 'courier' in fam) else FONT_SANS


def transform(text, tt):
    if tt == 'uppercase':
        return text.upper()
    if tt == 'lowercase':
        return text.lower()
    if tt == 'capitalize':
        return text.title()
    return text


def build_paragraphs(block):
    """Gộp run thành các đoạn (tách tại <br>), chuẩn hóa khoảng trắng như trình duyệt."""
    keep_ws = block['whiteSpace'] in ('pre', 'pre-wrap', 'pre-line', 'break-spaces')
    paras, cur = [], []
    for run in block['runs']:
        if run.get('br'):
            paras.append(cur)
            cur = []
            continue
        # Emoji/mũi tên đã được tách sang ảnh nền ở bước đo (slide_extract.js); chỉ bỏ ký tự điều khiển còn sót
        text = run['text'] if run.get('spacer') else run['text'].replace('\ufe0f', '').replace('\u200d', '')
        if not keep_ws and not run.get('spacer'):
            # Chỉ gộp khoảng trắng thường; giữ NBSP (chỗ của emoji trong ảnh nền)
            text = re.sub(r'[ \t\r\n\f]+', ' ', text)
        if text:
            cur.append((transform(text, run['style']['textTransform']), run['style']))
    paras.append(cur)

    cleaned = []
    for p in paras:
        if not keep_ws and p:
            p[0] = (p[0][0].lstrip(' '), p[0][1])
            p[-1] = (p[-1][0].rstrip(' '), p[-1][1])
        p = [(t, s) for t, s in p if t]
        cleaned.append(p)
    while cleaned and not cleaned[0]:
        cleaned.pop(0)
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return cleaned


ALIGN = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}


def add_text_block(slide, block, order):
    paras = build_paragraphs(block)
    if not paras:
        return False
    box = block['box']
    first_style = paras[0][0][1]
    font_px = first_style['fontSize']
    # Khoảng cách dòng: ưu tiên giá trị đo thực tế, sau đó line-height CSS,
    # cuối cùng là chiều cao dòng chữ (line-height: normal của font)
    line_px = box.get('pitch') or block['lineHeight'] or (box['b'] - box['t'])
    multi = box['lines'] > 1 or len(paras) > 1
    align = block['textAlign']

    if multi and block.get('display', '') in ('flex', 'inline-flex', 'grid', 'inline-grid'):
        # Chữ nằm trong ô flex/grid ẩn danh (hẹp hơn khung cha): dùng đúng phạm vi chữ
        left, right = box['l'], box['r'] + font_px * 0.5
    elif multi:
        # Nhiều dòng: dùng bề rộng vùng nội dung để xuống dòng giống bản gốc,
        # cộng 1.5% dư phòng chênh lệch nhỏ khi dàn chữ
        left, right = block['content']['l'], block['content']['r']
        left = min(left, box['l'])
        right = max(right, box['r'])
        extra = (right - left) * 0.015
        if align == 'center':
            left -= extra / 2
            right += extra / 2
        elif align in ('right', 'end'):
            left -= extra
        else:
            right += extra
    else:
        # 1 dòng: bề rộng chữ + 12% dư phòng chênh lệch font, không cho tự xuống dòng
        width = (box['r'] - box['l'])
        slack = width * 0.12 + font_px
        left, right = box['l'], box['r'] + slack
        if align == 'center':
            left -= slack / 2
            right -= slack / 2
        elif align in ('right', 'end'):
            left -= slack
            right -= slack

    # Đặt đỉnh hộp theo dòng chữ (bù nửa khoảng cách dòng)
    content_h = box['b'] - box['t']
    lines = max(box['lines'], len(paras))
    half_leading = max(0.0, (line_px * lines - content_h) / (2 * lines))
    top = box['t'] - half_leading
    height = max(line_px * lines, content_h)

    shape = slide.shapes.add_textbox(px(left), px(top), px(max(right - left, 4)), px(height))
    shape.name = f'Chữ {order}'
    tf = shape.text_frame
    tf.word_wrap = multi
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
        p.line_spacing = Pt(line_px * PT_PER_PX)
        p.space_before = p.space_after = Pt(0)
        if not para:
            continue
        for text, st in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = font_for(st['fontFamily'])
            # Làm tròn XUỐNG 0.1pt: chữ trên Canva chỉ có thể hẹp hơn bản đo, không bị xuống dòng thừa
            f.size = Pt(math.floor(st['fontSize'] * PT_PER_PX * 10) / 10)
            f.bold = st['fontWeight'] >= 600
            f.italic = st['italic']
            f.underline = st['underline']
            c = st['color']
            f.color.rgb = RGBColor(int(c['r']), int(c['g']), int(c['b']))
            if st['letterSpacing']:
                r._r.get_or_add_rPr().set('spc', str(int(round(st['letterSpacing'] * PT_PER_PX * 100))))
    return True


def build_pptx(slides, layouts, backgrounds, work_dir):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    stats = []
    for (num, notes, has_qr), layout, bg in zip(slides, layouts, backgrounds):
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(bg, 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = f'Nền slide {num}'

        for i, info in enumerate(layout['images']):
            fit = info.get('objectFit')
            if fit in ('contain', 'scale-down') and info['naturalW'] and info['naturalH']:
                # Ảnh được thu vừa khung, giữ tỉ lệ: đặt đúng vùng ảnh hiển thị (căn giữa)
                scale = min(info['w'] / info['naturalW'], info['h'] / info['naturalH'])
                if fit == 'scale-down':
                    scale = min(scale, 1)
                w, h = info['naturalW'] * scale, info['naturalH'] * scale
                info = {**info, 'l': info['l'] + (info['w'] - w) / 2, 't': info['t'] + (info['h'] - h) / 2,
                        'w': w, 'h': h, 'objectFit': 'fill'}
            path = prepare_image(info, work_dir, f'{num}_{i}')
            im = slide.shapes.add_picture(path, px(info['l']), px(info['t']), px(info['w']), px(info['h']))
            im.name = f'Ảnh {i + 1}'
            if 'qr_web_app' in info['src']:
                im.click_action.hyperlink.address = APP_URL

        count = sum(add_text_block(slide, b, i + 1) for i, b in enumerate(layout['blocks']))
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        stats.append((num, count, len(layout['images'])))

    prs.core_properties.title = 'Lừa đảo trực tuyến nhắm vào sinh viên — Tổ 4'
    prs.core_properties.author = 'Tổ 4 • Tư duy hệ thống'
    prs.save(OUT_PPTX)
    return stats


def main():
    slides = read_slides()
    if not slides:
        sys.exit('Không đọc được slide nào trong slides/index.html')
    only = {int(a) for a in sys.argv[1:] if a.isdigit()}
    if only:
        slides = [s for s in slides if s[0] in only]

    browser = find_browser()
    wrapper = os.path.join(ROOT, WRAPPER_NAME)
    server, port = start_server()
    print(f'Dựng PPTX chỉnh sửa được cho {len(slides)} slide bằng {os.path.basename(browser)} ...')
    try:
        with tempfile.TemporaryDirectory() as work_dir, tempfile.TemporaryDirectory() as profile_dir:
            with open(wrapper, 'w', encoding='utf-8') as fh:
                fh.write(wrapper_html('', ''))
            layouts, backgrounds = [], []
            for num, _, _ in slides:
                layouts.append(extract_slide(browser, port, profile_dir, num))
                backgrounds.append(screenshot_background(browser, port, profile_dir, work_dir, num))
                print(f'  ✓ Slide {num:2d}: {len(layouts[-1]["blocks"])} khối chữ, {len(layouts[-1]["images"])} ảnh')
            stats = build_pptx(slides, layouts, backgrounds, work_dir)
    finally:
        server.shutdown()
        if os.path.exists(wrapper):
            os.remove(wrapper)
    total_text = sum(s[1] for s in stats)
    print(f'→ {os.path.relpath(OUT_PPTX, ROOT)} ({os.path.getsize(OUT_PPTX) / 1e6:.1f} MB, {total_text} hộp chữ)')


if __name__ == '__main__':
    main()
