"""Xuất bộ slide Bản 4 (slides/index.html) ra PDF và PPTX chuẩn 16:9.

Cách dùng (từ thư mục gốc dự án):
    python tools/export_slides.py

Yêu cầu: Microsoft Edge (có sẵn trên Windows), python-pptx, Pillow.
Kết quả:
    slides/Slide_To4_DeTai9.pdf   — PDF 16:9, mỗi trang 1 slide (ảnh nét 2x)
    slides/Slide_To4_DeTai9.pptx  — PowerPoint 16:9 (13.333 x 7.5 in), kèm ghi chú người nói
"""

import functools
import html
import http.server
import os
import re
import shutil
import socketserver
import subprocess
import sys
import tempfile
import threading

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SLIDES_DIR = os.path.join(ROOT, 'slides')
SLIDES_HTML = os.path.join(SLIDES_DIR, 'index.html')
OUT_PDF = os.path.join(SLIDES_DIR, 'Slide_To4_DeTai9.pdf')
OUT_PPTX = os.path.join(SLIDES_DIR, 'Slide_To4_DeTai9.pptx')
APP_URL = 'https://chude9to4.vercel.app/app/index.html'

SLIDE_W, SLIDE_H = 1600, 900   # kích thước CSS của 1 slide
SCALE = 2                      # chụp 3200x1800 cho chữ sắc nét khi trình chiếu
JPEG_QUALITY = 92

EDGE_PATHS = [
    r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
    r'C:\Program Files\Microsoft\Edge\Application\msedge.exe',
]


def find_browser():
    for p in EDGE_PATHS:
        if os.path.exists(p):
            return p
    for name in ('msedge', 'chrome', 'google-chrome', 'chromium'):
        found = shutil.which(name)
        if found:
            return found
    sys.exit('Không tìm thấy Microsoft Edge / Chrome để render slide.')


def start_server():
    """Phục vụ thư mục gốc qua HTTP để font/ảnh tải giống khi mở trên web."""
    handler = functools.partial(QuietHandler, directory=ROOT)
    server = socketserver.ThreadingTCPServer(('127.0.0.1', 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, server.server_address[1]


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def read_slides():
    """Trả về danh sách (số slide, ghi chú người nói, có mã QR hay không)."""
    src = open(SLIDES_HTML, encoding='utf-8').read()
    sections = re.findall(r'<section class="slide-item[^"]*" data-slide="(\d+)"(.*?)</section>', src, re.S)
    slides = []
    for num, body in sections:
        notes_match = re.search(r'<div class="speaker-notes-content"[^>]*>(.*?)</div>', body, re.S)
        notes = ''
        if notes_match:
            notes = re.sub(r'<br\s*/?>', '\n', notes_match.group(1))
            notes = html.unescape(re.sub(r'<[^>]+>', '', notes))
            notes = '\n'.join(line.strip() for line in notes.splitlines() if line.strip())
        slides.append((int(num), notes, 'qr_web_app.png' in body))
    return slides


def run_browser(browser, profile_dir, args):
    cmd = [
        browser,
        '--headless=new',
        '--disable-gpu',
        '--hide-scrollbars',
        '--no-first-run',
        '--no-default-browser-check',
        f'--user-data-dir={profile_dir}',
        '--virtual-time-budget=8000',
        *args,
    ]
    subprocess.run(cmd, capture_output=True, text=True, timeout=120)


def screenshot_slides(browser, port, slides, work_dir, profile_dir):
    images = []
    for num, _, _ in slides:
        png = os.path.join(work_dir, f'slide_{num:02d}.png')
        url = f'http://127.0.0.1:{port}/slides/index.html?export=1&slide={num}'
        run_browser(browser, profile_dir, [
            f'--window-size={SLIDE_W},{SLIDE_H}',
            f'--force-device-scale-factor={SCALE}',
            f'--screenshot={png}',
            url,
        ])
        if not os.path.exists(png):
            sys.exit(f'Chụp slide {num} thất bại.')
        img = Image.open(png).convert('RGB')
        expected = (SLIDE_W * SCALE, SLIDE_H * SCALE)
        if img.size != expected:
            img = img.crop((0, 0, min(img.width, expected[0]), min(img.height, expected[1]))).resize(expected, Image.LANCZOS)
        jpg = os.path.join(work_dir, f'slide_{num:02d}.jpg')
        img.save(jpg, 'JPEG', quality=JPEG_QUALITY, subsampling=0, optimize=True)
        images.append(jpg)
        print(f'  ✓ Slide {num:2d}/{len(slides)}')
    return images


def build_pptx(slides, images):
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in — chuẩn 16:9 của PowerPoint
    prs.slide_height = Emu(6858000)   # 7.5 in
    blank = prs.slide_layouts[6]
    for (num, notes, has_qr), img in zip(slides, images):
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = f'Slide {num}'
        if has_qr:
            pic.click_action.hyperlink.address = APP_URL
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.core_properties.title = 'Lừa đảo trực tuyến nhắm vào sinh viên — Tổ 4'
    prs.core_properties.author = 'Tổ 4 • Tư duy hệ thống'
    prs.save(OUT_PPTX)


def build_pdf(images):
    pages = [Image.open(p).convert('RGB') for p in images]
    # 1600x900 CSS px ở 96 dpi → trang 16.667 x 9.375 in (16:9)
    dpi = 96 * SCALE
    pages[0].save(OUT_PDF, 'PDF', save_all=True, append_images=pages[1:], resolution=dpi,
                  title='Lừa đảo trực tuyến nhắm vào sinh viên — Tổ 4', author='Tổ 4')


def main():
    slides = read_slides()
    if not slides:
        sys.exit('Không đọc được slide nào trong slides/index.html')
    browser = find_browser()
    server, port = start_server()
    print(f'Render {len(slides)} slide bằng {os.path.basename(browser)} ...')
    with tempfile.TemporaryDirectory() as work_dir, tempfile.TemporaryDirectory() as profile_dir:
        try:
            images = screenshot_slides(browser, port, slides, work_dir, profile_dir)
        finally:
            server.shutdown()
        build_pptx(slides, images)
        build_pdf(images)
    for path in (OUT_PDF, OUT_PPTX):
        print(f'→ {os.path.relpath(path, ROOT)} ({os.path.getsize(path) / 1e6:.1f} MB)')


if __name__ == '__main__':
    main()
